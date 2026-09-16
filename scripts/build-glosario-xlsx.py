# -*- coding: utf-8 -*-
import json, re, collections, sys, os, unicodedata
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib.util as _u
_sp=_u.spec_from_file_location('gdefs', os.path.join(os.path.dirname(os.path.abspath(__file__)),'glosario-definiciones.py'))
_m=_u.module_from_spec(_sp); _sp.loader.exec_module(_m); DEFS=_m.DEFS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BASE=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
d=json.load(open(BASE+'/lib/prestaciones.json')); items=d['items']
tok=collections.Counter()
for i in items:
    for w in re.findall(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{2,}", i['n']): tok[w.upper()]+=1
fam=json.load(open(BASE+'/datos/glosario/terminos.json'))
de_fam={}
for k,v in fam.items():
    if k.startswith('—') or k.startswith('SIN') or k.startswith('Transparente'): continue
    for w in v: de_fam[w]=k
CAT={'lab':'Laboratorio','cir':'Cirugía','img':'Imagen','esp':'Especialidad','exc':'Exclusión','fis':'Fisioterapia'}

NAVY='FF003B71'; MINT='FFE6F7F6'; GOLD='FFF8F1DE'; GREY='FFF5F5F5'
H=Font(name='Arial',size=10,bold=True,color='FFFFFFFF')
B=Font(name='Arial',size=10,bold=True); N=Font(name='Arial',size=10)
S=Font(name='Arial',size=9,color='FF6B6B6B')
TIT=Font(name='Arial',size=14,bold=True,color=NAVY)
fillH=PatternFill('solid',fgColor=NAVY); fillM=PatternFill('solid',fgColor=MINT)
fillG=PatternFill('solid',fgColor=GOLD); fillGr=PatternFill('solid',fgColor=GREY)
wrap=Alignment(wrap_text=True,vertical='top'); top=Alignment(vertical='top')
thin=Border(bottom=Side(style='thin',color='FFDDDDDD'))

wb=Workbook(); wb.remove(wb.active)

# ─────────────────────────────── 1 · Cómo revisar
ws=wb.create_sheet('Cómo revisar')
ws.column_dimensions['A'].width=3; ws.column_dimensions['B'].width=104
def line(r,txt,f=N,fill=None):
    c=ws.cell(row=r,column=2,value=txt); c.font=f; c.alignment=wrap
    if fill: c.fill=fill
    return r+1
r=2
r=line(r,'Glosario médico de Salud Protegida — borradores para revisión',TIT); r+=1
r=line(r,'Qué es esto: los términos médicos que aparecen en las 983 coberturas publicadas en "¿Está cubierto lo '
        'que me pidieron?", con una definición borrador escrita en idioma de familia. NINGUNA se publica sin que '
        'un médico la valide.'); r+=1
r=line(r,'⚠ Los borradores los escribió Claude, no un médico. Son un punto de partida para corregir, no una '
        'fuente. Si una definición está mal, corregila en la columna "Corrección"; si está bien, marcá "Sí".',N,fillG); r+=1
r=line(r,'Cómo llenar cada fila',B)
r=line(r,'· Columna ¿Correcta?  — elegí Sí / No / Sacar del glosario.')
r=line(r,'· Columna Corrección  — si pusiste No, escribí acá cómo debería decir. Si pusiste Sí, dejala vacía.')
r=line(r,'· Columnas Revisor y Fecha — quién la firmó y cuándo. Sin eso no se publica.'); r+=1
r=line(r,'La vara de la definición',B)
r=line(r,'· Que la entienda alguien que no sabe nada de medicina. Si una definición necesita otra palabra difícil '
        'para explicarse, está mal escrita.')
r=line(r,'· Que diga PARA QUÉ SIRVE, no qué quiere decir la raíz griega.')
r=line(r,'· Sin prometer ni asustar. La definición explica, no vende ni diagnostica.'); r+=1
r=line(r,'Los tres estados de la columna Tipo',B)
r=line(r,'· Definición  — término real, con borrador escrito. Es el grueso.')
r=line(r,'· Modificador — palabra que solo tiene sentido dentro de la fila ("abierto", "por vena"). No va al '
        'glosario; se explica en la ficha.')
r=line(r,'· A confirmar — no sabemos qué significa en ESTA grilla. Necesitamos que Lister nos diga.'); r+=2
r=line(r,'Avance de la revisión',B)
fila_res=r
for et,form in [('Total de términos','=COUNTA(Términos!A2:A400)'),
                ('Revisados','=COUNTIF(Términos!F2:F400,"<>")'),
                ('Marcados Sí','=COUNTIF(Términos!F2:F400,"Sí")'),
                ('Marcados No','=COUNTIF(Términos!F2:F400,"No")'),
                ('Pendientes','=COUNTA(Términos!A2:A400)-COUNTIF(Términos!F2:F400,"<>")')]:
    ws.cell(row=r,column=2,value=et).font=N
    c=ws.cell(row=r,column=3,value=form); c.font=B; c.alignment=top
    r+=1
ws.column_dimensions['C'].width=12

# ─────────────────────────────── 2 · Términos
ws=wb.create_sheet('Términos')
cols=[('Término',20),('Familia',34),('Tipo',13),('Fichas',8),('Definición borrador (corregir acá al lado)',62),
      ('¿Correcta?',13),('Corrección',52),('Revisor',14),('Fecha',12)]
for j,(t,w) in enumerate(cols,1):
    c=ws.cell(row=1,column=j,value=t); c.font=H; c.fill=fillH; c.alignment=wrap
    ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='A2'
def canon(w): return ''.join(c for c in unicodedata.normalize('NFD',w) if unicodedata.category(c)!='Mn')
vistos={}
for w,(tipo,txt) in DEFS.items():
    k=canon(w)
    if k in vistos: continue
    vistos[k]=(w,tipo,txt)
filas=[]
for w,(tipo,txt) in DEFS.items():
    if vistos.get(canon(w),(None,))[0]!=w: continue
    variantes=[x for x in tok if canon(x)==canon(w)]
    filas.append((w, de_fam.get(w,'—'), {'D':'Definición','M':'Modificador','C':'A confirmar'}[tipo],
                  sum(tok[x] for x in variantes), txt))
filas.sort(key=lambda x:(x[1], -x[3], x[0]))
TIPO_FILL={'Modificador':fillGr,'A confirmar':fillG}
for i,(t,f,tp,n,txt) in enumerate(filas,2):
    for j,v in enumerate([t,f,tp,n,txt],1):
        c=ws.cell(row=i,column=j,value=v); c.font=B if j==1 else N
        c.alignment=wrap if j==5 else top; c.border=thin
        if j==3 and tp in TIPO_FILL: c.fill=TIPO_FILL[tp]
    for j in range(6,10):
        ws.cell(row=i,column=j).fill=fillM; ws.cell(row=i,column=j).border=thin
dv=DataValidation(type='list',formula1='"Sí,No,Sacar del glosario"',allow_blank=True)
ws.add_data_validation(dv); dv.add(f'F2:F{len(filas)+1}')
n_term=len(filas)

# ─────────────────────────────── 3 · Fichas que quedan sin cubrir
cubiertas=[]; sin=[]
dicc={canon(w) for w in DEFS}
for i in items:
    ws_=[canon(x.upper()) for x in re.findall(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{2,}", i['n'])]
    (cubiertas if dicc & set(ws_) else sin).append(i)
ws=wb.create_sheet('Fichas sin cubrir')
for j,(t,w) in enumerate([('Nombre en el tarifario',62),('Cuadro',14),
                          ('Explicación de una línea (escribir acá)',62),('Revisor',14),('Fecha',12)],1):
    c=ws.cell(row=1,column=j,value=t); c.font=H; c.fill=fillH; c.alignment=wrap
    ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='A2'
for i,it in enumerate(sorted(sin,key=lambda x:(x.get('c',''),x['n'])),2):
    ws.cell(row=i,column=1,value=it['n']).font=N
    ws.cell(row=i,column=2,value=CAT.get(it.get('c'),it.get('c',''))).font=N
    for j in range(1,6):
        ws.cell(row=i,column=j).alignment=wrap if j in (1,3) else top
        ws.cell(row=i,column=j).border=thin
        if j>=3: ws.cell(row=i,column=j).fill=fillM
wb['Cómo revisar'].cell(row=fila_res-2,column=2,value=
    f'Esta primera tanda son {n_term} términos. Con ellos, {len(cubiertas)} de las {len(items)} coberturas quedan con al menos '
    f'una palabra explicada. Las {len(sin)} restantes están en la hoja "Fichas sin cubrir": esas no se arreglan con '
    'una palabra, necesitan una línea propia — es la tanda 2.').font=N
wb['Cómo revisar'].cell(row=fila_res-2,column=2).alignment=wrap

out=BASE+'/datos/glosario/Glosario-medico-SP-borradores.xlsx'
wb.save(out)
print(f'✓ {out}')
print(f'  Términos: {n_term} · fichas cubiertas: {len(cubiertas)}/{len(items)} · sin cubrir: {len(sin)}')
