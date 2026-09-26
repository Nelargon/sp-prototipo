# -*- coding: utf-8 -*-
"""Genera lib/guia-medica.json desde la planilla maestra de la Guía Médica.

La planilla (Excel) es la fuente: la arma y la mantiene SP, y vive en el repo
PRIVADO (sp-interno/project/guia-medica/) porque trae Observaciones internas
("confirmar con Tesorería", "confirmar si cambió de sede"). Acá sale solo lo
que la guía online muestra. Es INTERINO (decisión de Arturo, 23/09/2026):
el destino final es leer la red directo del sistema de SP.

Uso (openpyxl vive FUERA del repo, como playwright — regla del CLAUDE.md):
    python3 -m venv /tmp/venv && /tmp/venv/bin/pip install openpyxl
    /tmp/venv/bin/python scripts/build-guia-medica.py <planilla.xlsx>                 # valida, informa y escribe
    /tmp/venv/bin/python scripts/build-guia-medica.py <planilla.xlsx> --solo-validar  # valida e informa, no escribe
    ... --informe informe.md   # además guarda el informe (va en la descripción del PR)

Control antes de publicar (23/09/2026, acordado con Cowork):
1. VALIDA la planilla. Los ERRORES frenan: no se escribe nada y sale con código
   1. Los AVISOS no frenan. Los controles son los de `validar_y_exportar_guia.py`
   (Cowork), portados todos — ver `validar()`.
2. INFORMA qué cambió contra la última publicación (el `lib/guia-medica.json`
   que está en git): altas, bajas y cambios de teléfono, dirección, nombre,
   especialidad, ciudad, redes y estado «Revisar». Y dice qué reglas aplicó.
3. PUBLICA recién con el OK de Arturo: el PR de datos de la guía NO se fusiona
   solo (excepción al merge automático del CLAUDE.md). El informe va en el PR.

La planilla original vive en el Drive de Arturo (ver sp-interno/README.md: se
edita solo esa). La copia de sp-interno es la foto de la última publicación.

Qué decide este script (y por qué):
- Redes. Cada guía en PDF es una red. Se publican las que tienen planes
  conocidos: Privilege y su familia, Essential en sus tres zonas (plan
  vigente; NO es el "SP Esencial" de la gama nueva, que todavía no existe —
  Arturo, 25/09/2026, corrigiendo lo anotado el 23/09), y el Seguro Estatal. La
  Centralizada NO: el PDF no dice qué planes cubre, y una red que no se puede
  asociar a un plan no le contesta nada a nadie. Las filas que solo están en
  la Centralizada quedan afuera.
- Estado. "Baja" no se publica. "Revisar" se publica igual que en el PDF y
  lleva `revisar: true` (el prototipo lo marca con un punto naranja, para
  revisión interna; la v1 no lo muestra).
- Tipo "Aviso" no es un prestador: la guía pone una leyenda en lugar de una
  lista. Viaja como nota de la especialidad.
- Condiciones: se quita la procedencia "(según …)", que nombra guías que la
  persona no conoce. Y la condición del Plan Centralizado («Plan
  Centralizado: …») no se publica: esa guía no está online (Decisiones
  23/09, n.º 4), y un asegurado de otro plan la leería como propia (26/09).
"""
import sys, os, re, json, unicodedata, argparse, subprocess, datetime
from collections import Counter, defaultdict
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(BASE, 'lib', 'guia-medica.json')

# columna de la planilla → clave de red en la web
REDES = [
    ('Privilege y otros', 'privilege'),
    ('Essential Asu-Central', 'esencial_ac'),
    ('Essential Interior', 'esencial_int'),
    ('Essential Nacional', 'esencial_nac'),
    ('Estatal V2', 'estatal'),
]
GUIA_DE_RED = {'privilege': 'privilege', 'esencial_ac': 'ess_asucentral',
               'esencial_int': 'ess_interior', 'esencial_nac': 'ess_nacional',
               'estatal': 'estatal'}


def txt(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def sin_segun(c):
    partes = [re.sub(r'\s*\(según [^)]*\)', '', p).strip() for p in c.split(';')]
    vistos, out = set(), []
    for p in partes:
        if p and p.lower() not in vistos:
            vistos.add(p.lower()); out.append(p)
    return '; '.join(out)


CONDS_CENTRAL = []  # filas a las que se les sacó la condición de la Centralizada


def sin_centralizada(c):
    return '; '.join(p for p in c.split('; ') if p and not p.lower().startswith('plan centralizado'))


def telefonos(*vals):
    out = []
    for v in vals:
        for t in re.split(r'\s*/\s*', txt(v)):
            if t and t not in out:
                out.append(t)
    return out


# Lister es Pa'i Pérez 630. Buscar solo la palabra "lister" no alcanzaba:
# cinco filas en esa misma dirección dicen "C.M. Salud Protegida" y quedaban
# sin la marca de centro propio (hallazgo del code review, 23/09/2026). Vale
# cualquiera de tres señales: la palabra, la dirección, o que la hoja «Lister»
# de la planilla nombre la fila.
LISTER_DIR = re.compile(r"pa\W?i\s*p[eé]rez\s*(n[º°o]\.?\s*)?630\b", re.I)
FILAS_LISTER = set()


def en_lister(direccion, nombre, id_fila):
    return ('lister' in (direccion + ' ' + nombre).lower()
            or bool(LISTER_DIR.search(direccion))
            or id_fila in FILAS_LISTER)


def clave_orden(nombre):
    n = re.sub(r'^(Dra?\.|Lic\.|Prof\.|Od\.)\s*', '', nombre)
    return ''.join(c for c in unicodedata.normalize('NFD', n.lower()) if unicodedata.category(c) != 'Mn')


# ---------------------------------------------------------------- validación
# Portado de `validar_y_exportar_guia.py` (Cowork, 23/09/2026), control por
# control. Busca las columnas por su título, no por su posición: si alguien
# agrega o mueve una columna, la validación no se corre en silencio de lugar.
COLS = ['ID fila', 'Estado', 'Especialidad', 'Nombre', 'Tipo', 'Dirección', 'Barrio', 'Ciudad', 'Departamento',
        'Teléfono', 'Otros teléfonos', 'Condiciones', 'Observaciones', 'Fuente (guía y página)', 'ID prestador']
ESTADOS = {'Activo', 'Revisar', 'Baja'}
TIPOS = {'Profesional', 'Institución', 'Aviso'}
TEL = re.compile(r'^\+595 (9\d\d \d{3} \d{3}|21 \d{3} (\d{3}|\d{2} \d{2})|\d{2,3} (\d{2} \d{3}|\d{3} \d{3}|\d{3} \d{2} \d{2}))$')


def K(s):
    s = unicodedata.normalize('NFKD', str(s or ''))
    return re.sub(r'\s+', ' ', ''.join(c for c in s if not unicodedata.combining(c)).upper()).strip()


def leer_catalogos(wb):
    cat = list(wb['Catálogos'].iter_rows(values_only=True))
    hdr = [txt(h) for h in cat[0]]

    def col(nombre, n=1):
        idx = [i for i, h in enumerate(hdr) if h == nombre]
        if len(idx) < n:
            sys.exit(f'✗ En Catálogos falta la columna «{nombre}».')
        return idx[n - 1]
    c_esp, c_gru, c_dep = col('Especialidad'), col('Grupo (guía online)'), col('Departamento', 1)
    c_ciu, c_cdep, c_bar = col('Ciudad'), col('Departamento', 2), col('Barrio (Asunción)')
    c_gid, c_gnom, c_gfec, c_gpla = col('Guía'), col('Nombre'), col('Fecha del PDF'), col('Planes')
    esp, deptos, ciudades, barrios, guias = {}, set(), {}, set(), {}
    for r in cat[1:]:
        r = list(r) + [None] * (len(hdr) - len(r))
        if txt(r[c_esp]): esp[txt(r[c_esp])] = txt(r[c_gru])
        if txt(r[c_dep]): deptos.add(txt(r[c_dep]))
        if txt(r[c_ciu]): ciudades[txt(r[c_ciu])] = txt(r[c_cdep])
        if txt(r[c_bar]): barrios.add(txt(r[c_bar]))
        if txt(r[c_gid]):
            guias[txt(r[c_gid])] = {'nombre': txt(r[c_gnom]), 'fecha': txt(r[c_gfec]), 'planes': txt(r[c_gpla])}
    return esp, deptos, ciudades, barrios, guias


def validar(red, n_fila, esp, deptos, ciudades, barrios, guias):
    """Devuelve (errores, avisos). Errores frenan la publicación."""
    err, av = [], []
    gcols = {g['nombre']: gid for gid, g in guias.items()}
    ids = Counter(txt(r['ID fila']) for r in red)
    vistos = defaultdict(list)
    for r in red:
        f = {c: txt(r.get(c)) for c in COLS}
        n, fid = n_fila[id(r)], f['ID fila'] or f'(fila {n_fila[id(r)]})'
        E = lambda m: err.append(f'Fila {n} · {fid}: {m}')
        A = lambda m: av.append(f'Fila {n} · {fid}: {m}')
        # ID de FILA único; el ID de PRESTADOR se repite a propósito (otra sede
        # u otra especialidad del mismo prestador).
        if not re.fullmatch(r'F-\d{4,}', f['ID fila']): E('ID fila vacío o con formato distinto de F-0000')
        elif ids[f['ID fila']] > 1: E('ID fila repetido')
        if not re.fullmatch(r'P-\d{4,}', f['ID prestador']): E('ID prestador vacío o con formato distinto de P-0000')
        if f['Estado'] not in ESTADOS: E(f'Estado «{f["Estado"]}» no válido (Activo, Revisar o Baja)')
        if f['Tipo'] not in TIPOS: E(f'Tipo «{f["Tipo"]}» no válido')
        if f['Especialidad'] not in esp: E(f'Especialidad «{f["Especialidad"]}» no está en Catálogos')
        if not f['Nombre']: E('Nombre vacío')
        marcas = {g: txt(r.get(col)) for col, g in gcols.items()}
        malas = [g for g, v in marcas.items() if v not in ('', 'Sí')]
        if malas: E('En las columnas de guías solo va «Sí» o nada: ' + ', '.join(malas))
        if f['Estado'] != 'Baja' and 'Sí' not in marcas.values(): E('No está marcada en ninguna guía')
        # Un «Aviso» es una leyenda de la guía, no un prestador: no lleva
        # dirección ni teléfono, y no frena.
        if f['Tipo'] != 'Aviso':
            for c in ('Dirección', 'Ciudad', 'Departamento'):
                if not f[c]: E(f'{c} vacía')
            if not f['Teléfono']: A('Sin teléfono')
        if f['Departamento'] and f['Departamento'] not in deptos: E(f'Departamento «{f["Departamento"]}» no válido')
        if f['Ciudad']:
            if f['Ciudad'] not in ciudades: A(f'Ciudad «{f["Ciudad"]}» no está en Catálogos (sumala con su departamento)')
            elif f['Departamento'] and ciudades[f['Ciudad']] != f['Departamento']:
                E(f'{f["Ciudad"]} es de {ciudades[f["Ciudad"]]}, no de {f["Departamento"]}')
        if f['Barrio']:
            if f['Ciudad'] != 'Asunción': A('Tiene barrio pero la ciudad no es Asunción')
            elif f['Barrio'] not in barrios: A(f'Barrio «{f["Barrio"]}» no está en Catálogos')
        for t in [t for t in [f['Teléfono']] + [x.strip() for x in f['Otros teléfonos'].split('/')] if t]:
            if not TEL.match(t): E(f'Teléfono «{t}» con formato no válido (ej.: +595 21 319 00 00 · +595 981 123 456)')
        if f['Tipo'] == 'Profesional' and not re.match(r'(Dra?|Lic)\.\s', f['Nombre']): A('Profesional sin título (Dr., Dra., Lic.)')
        if f['Estado'] != 'Baja':
            vistos[(f['Especialidad'], K(re.sub(r'^(Dra?|Lic)\.\s*', '', f['Nombre'])), K(f['Dirección']))].append(fid)
    for k, v in vistos.items():
        if len(v) > 1: av.append(f'Posible repetido ({k[0]} · misma persona y dirección): ' + ', '.join(v))
    return err, av


# ------------------------------------------------ informe de cambios
CAMPOS = [('n', 'nombre'), ('e', 'especialidad'), ('tel', 'teléfono'), ('d', 'dirección'),
          ('c', 'ciudad'), ('r', 'redes'), ('rv', 'Revisar'), ('k', 'condiciones')]


def publicado():
    """El JSON de la última publicación: el que está en git (HEAD)."""
    try:
        crudo = subprocess.check_output(['git', 'show', 'HEAD:lib/guia-medica.json'], cwd=BASE, stderr=subprocess.DEVNULL)
        return json.loads(crudo)
    except Exception:
        return None


def informe_cambios(antes, ahora):
    L = []
    if not antes:
        return ['No hay publicación anterior en git: todo es alta.']
    A = {p['f']: p for p in antes['prestadores']}
    B = {p['f']: p for p in ahora['prestadores']}
    altas, bajas = sorted(B.keys() - A.keys()), sorted(A.keys() - B.keys())
    cambios = []
    for f in sorted(A.keys() & B.keys()):
        for k, nom in CAMPOS:
            va, vb = A[f].get(k), B[f].get(k)
            if va != vb:
                fmt = lambda v: ('sí' if v == 1 else 'no') if k == 'rv' else (' / '.join(v) if isinstance(v, list) else (v or '—'))
                cambios.append(f'- {f} · {B[f]["n"]} · **{nom}**: {fmt(va)} → {fmt(vb)}')
    L.append(f'**{len(altas)} altas · {len(bajas)} bajas · {len(cambios)} cambios** '
             f'(contra la publicación anterior: {antes["meta"].get("generado_de", "?")}).')
    if altas:
        L.append('\n**Altas**')
        L += [f'- {f} · {B[f]["n"]} · {B[f]["e"]} · {B[f]["c"]} · {" / ".join(B[f]["tel"]) or "sin teléfono"}' for f in altas]
    if bajas:
        L.append('\n**Bajas** (dejan de verse online)')
        L += [f'- {f} · {A[f]["n"]} · {A[f]["e"]} · {A[f]["c"]}' for f in bajas]
    if cambios:
        L.append('\n**Cambios**')
        L += cambios
    na, nb = antes.get('notas', {}), ahora.get('notas', {})
    if na != nb:
        L.append('\n**Avisos de especialidad** cambiaron: ' + ', '.join(sorted(set(na) ^ set(nb) | {k for k in na.keys() & nb.keys() if na[k] != nb[k]})))
    if antes.get('lister') != ahora.get('lister'):
        L.append('\n**Horarios de Lister** cambiaron.')
    return L


def main(ruta, solo_validar=False, ruta_informe=None):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    for h in ('Red', 'Catálogos', 'Lister'):
        if h not in wb.sheetnames:
            sys.exit(f'✗ La planilla no tiene la hoja «{h}».')

    filas = list(wb['Red'].iter_rows(values_only=True))
    H = [txt(h) for h in filas[0]]
    faltan = [c for c in COLS + [col for col, _ in REDES] if c not in H]
    if faltan:
        sys.exit('✗ En Red faltan columnas: ' + ', '.join(faltan))
    red, n_fila = [], {}
    for n, f in enumerate(filas[1:], 2):
        if not any(txt(v) for v in f):
            continue
        r = dict(zip(H, f))
        n_fila[id(r)] = n
        red.append(r)

    grupo, deptos, ciudades, barrios, guias = leer_catalogos(wb)

    errores, avisos = validar(red, n_fila, grupo, deptos, ciudades, barrios, guias)
    est = Counter(txt(r['Estado']) for r in red)
    rep = [f'## Informe de la Guía Médica · {os.path.basename(ruta)}', '',
           f'{len(red)} filas en la planilla · Activo {est["Activo"]} · Revisar {est["Revisar"]} · Baja {est["Baja"]}', '',
           f'### Validación: {len(errores)} errores · {len(avisos)} avisos']
    rep += [f'- ✗ {m}' for m in errores] + [f'- · {m}' for m in avisos]
    if errores:
        rep += ['', '**No se publicó: corregí los errores en la planilla y volvé a correr.**']
        salir(rep, ruta_informe, 1)

    # Filas que la hoja «Lister» cruza con la Red (columna "Fila en la Red").
    for f in wb['Lister'].iter_rows(values_only=True):
        for c in f:
            if isinstance(c, str) and re.fullmatch(r'F-\d{4}', c.strip()):
                FILAS_LISTER.add(c.strip())

    prestadores, notas = [], {}
    excluidas = {'baja': 0, 'solo_centralizada': 0}
    for r in red:
        estado = txt(r['Estado'])
        if estado == 'Baja':
            excluidas['baja'] += 1; continue
        esp = txt(r['Especialidad'])
        if txt(r['Tipo']) == 'Aviso':
            notas[esp] = txt(r['Nombre']); continue
        redes = [k for col, k in REDES if txt(r[col]) == 'Sí']
        if not redes:
            excluidas['solo_centralizada'] += 1; continue
        direccion = txt(r['Dirección'])
        nombre = txt(r['Nombre'])
        p = {
            'f': txt(r['ID fila']),
            'id': txt(r['ID prestador']),
            'n': nombre,
            'e': esp,
            'g': grupo.get(esp, 'Especialidades médicas'),
            't': 'i' if txt(r['Tipo']) == 'Institución' else 'p',
            'd': direccion,
            'b': txt(r['Barrio']),
            'c': txt(r['Ciudad']),
            'dp': txt(r['Departamento']),
            'tel': telefonos(r['Teléfono'], r['Otros teléfonos']),
            'r': redes,
        }
        crudo = sin_segun(txt(r['Condiciones']))
        cond = sin_centralizada(crudo)
        if cond != crudo:
            CONDS_CENTRAL.append(txt(r['ID fila']))
        if cond:
            p['k'] = cond
        if en_lister(direccion, nombre, txt(r['ID fila'])):
            p['l'] = 1
        if estado == 'Revisar':
            p['rv'] = 1
        prestadores.append(p)

    prestadores.sort(key=lambda p: (p['e'], clave_orden(p['n'])))

    # Página «Guía Lister»: servicios y horarios (igual en las 6 guías)
    lister = []
    en_horarios = False
    for f in wb['Lister'].iter_rows(values_only=True):
        a, b = txt(f[0]), txt(f[1] if len(f) > 1 else '')
        if a == 'Servicio' and b.startswith('Horario'):
            en_horarios = True; continue
        if en_horarios:
            if not a or a.startswith('En la Red con'):
                break
            lister.append({'s': a, 'h': b})

    con_gente = {p['e'] for p in prestadores}
    sin_prest = sorted([e, g] for e, g in grupo.items() if e not in con_gente)
    out = {
        'meta': {
            'generado_de': re.sub(r'^[0-9a-f]{8}-', '', os.path.basename(ruta)),
            'datos_al': datetime.date.today().isoformat(),
            'guias': {GUIA_DE_RED[k]: guias.get(GUIA_DE_RED[k]) for _, k in REDES},
            'excluidas': excluidas,
            # Especialidades del catálogo que no quedaron con nadie publicado
            # (24/09/2026, pedido de Arturo): el prototipo las muestra con un
            # círculo naranja para decidir si se completan o se sacan. La v1
            # no las muestra.
            'sin_prestadores': sin_prest,
        },
        'notas': notas,
        'lister': lister,
        'prestadores': prestadores,
    }
    solo_central = [txt(r['ID fila']) for r in red if txt(r['Estado']) != 'Baja' and txt(r['Tipo']) != 'Aviso'
                    and not any(txt(r[c]) == 'Sí' for c, _ in REDES)]
    rep += ['', '### Reglas que se aplicaron',
            '- Guías publicadas: ' + ', '.join(f'{guias.get(GUIA_DE_RED[k], {}).get("nombre", k)} ({GUIA_DE_RED[k]})' for _, k in REDES) + '.',
            f'- Centralizada: fuera de la guía online por ahora (Decisiones 23/09, n.º 4). '
            f'{len(solo_central)} filas están solo ahí y no se publican: {", ".join(solo_central) or "ninguna"}.',
            f'- Condición «Plan Centralizado: …»: no se publica, por la misma razón '
            f'({len(CONDS_CENTRAL)} filas: {", ".join(CONDS_CENTRAL) or "ninguna"}).',
            f'- «Revisar»: se publican, como en el PDF ({sum(1 for p in prestadores if p.get("rv"))} filas). '
            'El punto naranja se ve solo en el prototipo, no en la v1.',
            f'- «Baja»: no se publica ({excluidas["baja"]} filas).',
            f'- Especialidades del catálogo sin nadie publicado ({len(sin_prest)}): '
            + (', '.join(e for e, _ in sin_prest) or 'ninguna') + '. En el prototipo llevan un círculo naranja; en la v1 no aparecen.',
            f'- Lister: hoja «Lister» de la planilla + dirección Pa\'i Pérez 630 ({sum(1 for p in prestadores if p.get("l"))} filas).',
            '', '### Qué cambia online'] + informe_cambios(publicado(), out)
    if solo_validar:
        rep += ['', '_Solo validación: no se escribió nada._']
        salir(rep, ruta_informe, 0)
    with open(SALIDA, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(',', ':'))
    # Resumen chico para el simulador: cuántos prestadores distintos tiene cada
    # red por departamento y ciudad. El simulador no carga la red entera (250
    # KB); le alcanza con esto para decir "en tu ciudad tenés N" y abrir la
    # guía ya filtrada.
    resumen = {}
    for p in prestadores:
        for r in p['r']:
            d = resumen.setdefault(r, {}).setdefault(p['dp'] or '—', {'ids': set(), 'c': {}})
            d['ids'].add(p['id'])
            d['c'].setdefault(p['c'], set()).add(p['id'])
    resumen = {r: {dp: {'n': len(v['ids']), 'c': {c: len(ids) for c, ids in v['c'].items()}} for dp, v in dps.items()}
               for r, dps in resumen.items()}
    with open(os.path.join(BASE, 'lib', 'red-resumen.json'), 'w', encoding='utf-8') as fh:
        json.dump(resumen, fh, ensure_ascii=False, separators=(',', ':'), sort_keys=True)

    n_prest = len({p['id'] for p in prestadores})
    rep += ['', f'✓ Escrito lib/guia-medica.json: {len(prestadores)} filas · {n_prest} prestadores. '
            '**Se publica recién con el OK de Arturo** (el PR no se fusiona solo).']
    salir(rep, ruta_informe, 0)


def salir(rep, ruta_informe, codigo):
    texto = '\n'.join(rep) + '\n'
    print(texto)
    if ruta_informe:
        with open(ruta_informe, 'w', encoding='utf-8') as fh:
            fh.write(texto)
    sys.exit(codigo)


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='Valida la planilla de la Guía Médica, informa qué cambia y genera lib/guia-medica.json.')
    ap.add_argument('planilla')
    ap.add_argument('--solo-validar', action='store_true', help='valida e informa, no escribe nada')
    ap.add_argument('--informe', help='guarda el informe en este archivo (markdown)')
    a = ap.parse_args()
    main(a.planilla, a.solo_validar, a.informe)
